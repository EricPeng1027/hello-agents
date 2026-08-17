"""read_data_table 数据表读取工具

作为 DraftingAgent 的本地工具，在 ReAct 撰写过程中可被调用，
读取用户上传的数据表（CSV/XLSX）并返回结构化数据（Markdown 表）。

支持用法：
- read_data_table[文件名]                    → 读全表（截断保护）
- read_data_table[文件名][关键词]            → 按关键词过滤行
- read_data_table[文件名][col=值]            → 按列名精确过滤
- read_data_table[文件名][col=值, 关键词]    → 组合过滤

文件解析约定：
- 输入只取文件名（basename），在 data/facts、data/style、data/ 下按序查找，
  防止绝对路径/路径穿越
- CSV 用标准库解析（自动嗅探 utf-8/utf-8-sig/gbk）；XLSX 依赖 openpyxl，
  未安装时返回明确提示而非报错中断 ReAct 循环
"""

import csv
from pathlib import Path
from typing import Callable, List, Optional

from ..materials.manager import SCOPE_FACTS, SCOPE_STYLE

# 数据表扩展名（上传白名单并入 loader.SUPPORTED_EXTENSIONS）
DATA_TABLE_EXTENSIONS = {".csv", ".xlsx"}

# 截断保护：避免超长表格撑爆 prompt
MAX_ROWS = 50
MAX_COLS = 12
MAX_CELL_CHARS = 80
MAX_TOTAL_CHARS = 3000


def build_read_data_table_tool(data_dir: str) -> Callable[[str], str]:
    """构建 read_data_table 工具函数

    Args:
        data_dir: 参考材料根目录（其下 facts/、style/ 子目录及各类型专属
            子目录 data_dir/<type_id>/facts|style 一并纳入查找）

    Returns:
        一个 (query: str) -> str 的可调用函数
    """
    root = Path(data_dir)

    def read_data_table(query: str) -> str:
        args = _split_args(query)
        if not args or not args[0]:
            return "（用法: read_data_table[文件名] 或 read_data_table[文件名][关键词]）"

        path = _locate(root, args[0])
        if path is None:
            return f"（未找到数据表 '{args[0]}'，请确认已上传到对应类型的 facts/style 库）"

        try:
            if path.suffix.lower() == ".csv":
                rows = _read_csv(path)
            else:
                rows = _read_xlsx(path)
        except ImportError as e:
            return f"（读取 {path.name} 失败: {e}）"
        except Exception as e:
            return f"（读取 {path.name} 失败: {e}）"

        if not rows:
            return f"（{path.name} 为空表或无法解析）"

        headers, data_rows = rows[0], rows[1:]
        filters = args[1:]
        if filters:
            data_rows = _filter_rows(headers, data_rows, filters)

        return _render_markdown(path.name, headers, data_rows)

    return read_data_table


# ------------------------------------------------------------------ 解析
def _split_args(query: str) -> List[str]:
    """拆分 '文件][条件1][条件2' 形式的参数（兼容单参数无第二组括号）"""
    parts = [p.strip() for p in query.split("][")]
    return [p.strip("[] ").strip() for p in parts if p.strip("[] ").strip()]


def _locate(root: Path, name: str) -> Optional[Path]:
    """按文件名查找数据表（只取 basename，防穿越）

    查找顺序：根下 facts/style → 各类型专属目录 <type>/facts|style → 根目录。
    """
    base = Path(name.strip()).name  # 剥离任何目录成分
    if not base:
        return None
    candidates = [root / SCOPE_FACTS / base, root / SCOPE_STYLE / base]
    if root.is_dir():
        for type_dir in sorted(p for p in root.iterdir() if p.is_dir()):
            candidates.append(type_dir / SCOPE_FACTS / base)
            candidates.append(type_dir / SCOPE_STYLE / base)
    candidates.append(root / base)
    for candidate in candidates:
        if (
            candidate.is_file()
            and candidate.suffix.lower() in DATA_TABLE_EXTENSIONS
        ):
            return candidate
    return None


def _read_csv(path: Path) -> List[List[str]]:
    """读取 CSV，自动尝试常见编码"""
    last_err: Optional[Exception] = None
    for enc in ("utf-8-sig", "utf-8", "gbk"):
        try:
            with open(path, "r", encoding=enc, newline="") as f:
                return [
                    [str(c).strip() for c in row] for row in csv.reader(f) if row
                ]
        except (UnicodeDecodeError, UnicodeError) as e:
            last_err = e
    raise ValueError(f"CSV 编码无法识别: {last_err}")


def _read_xlsx(path: Path) -> List[List[str]]:
    """读取 XLSX 第一个 sheet（依赖 openpyxl）"""
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ImportError("读取 .xlsx 需要安装 openpyxl（pip install openpyxl）")
    wb = load_workbook(str(path), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = []
    for row in ws.iter_rows(values_only=True):
        cells = ["" if c is None else str(c).strip() for c in row]
        if any(cells):
            rows.append(cells)
    wb.close()
    return rows


def _filter_rows(
    headers: List[str], rows: List[List[str]], filters: List[str]
) -> List[List[str]]:
    """按 col=值（精确）与裸关键词（任意列包含）组合过滤"""
    result = rows
    for cond in filters:
        if "=" in cond:
            col, _, val = cond.partition("=")
            col, val = col.strip(), val.strip()
            try:
                idx = [h.strip() for h in headers].index(col)
            except ValueError:
                return [f"（列 '{col}' 不存在，可用列: {', '.join(headers)}）"]
            result = [r for r in result if idx < len(r) and r[idx] == val]
        else:
            kw = cond.strip()
            if kw:
                result = [r for r in result if any(kw in c for c in r)]
    return result


def _render_markdown(name: str, headers: List[str], rows: List[List[str]]) -> str:
    """渲染为截断保护的 Markdown 表"""
    truncated: List[str] = []
    total_rows = len(rows)

    headers = [h[:MAX_CELL_CHARS] for h in headers[:MAX_COLS]]
    if len(headers) < 1:
        return f"（{name} 无表头）"

    out_rows = []
    for r in rows[:MAX_ROWS]:
        cells = [(c[:MAX_CELL_CHARS] if c else "") for c in r[:MAX_COLS]]
        cells += [""] * (len(headers) - len(cells))
        out_rows.append(cells)
    if total_rows > MAX_ROWS:
        truncated.append(f"行数 {total_rows}，仅展示前 {MAX_ROWS} 行")

    lines = [f"【数据表: {name}】", ""]
    lines.append("| " + " | ".join(headers) + " |")
    lines.append("| " + " | ".join("---" for _ in headers) + " |")
    for cells in out_rows:
        lines.append("| " + " | ".join(cells) + " |")

    text = "\n".join(lines)
    if len(text) > MAX_TOTAL_CHARS:
        text = text[:MAX_TOTAL_CHARS] + "\n...（内容过长已截断）"
    if truncated:
        text += "\n（" + "；".join(truncated) + "）"
    return text


# 工具元信息（供注册时使用）
TOOL_NAME = "read_data_table"
TOOL_DESCRIPTION = (
    "读取用户上传的数据表（CSV/XLSX），返回 Markdown 表格数据。"
    "格式: read_data_table[文件名] 读全表；"
    "read_data_table[文件名][关键词] 按关键词过滤行；"
    "read_data_table[文件名][列名=值] 按列精确过滤。"
    "撰写数据型章节（关键数据/量化指标）时优先用它取真实数据。"
)
